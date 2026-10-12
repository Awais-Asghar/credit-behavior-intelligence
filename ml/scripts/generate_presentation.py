"""Generate the editable Customer Intelligence portfolio presentation."""

from __future__ import annotations

import argparse
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
FIGURES = ROOT / "reports" / "figures"
ASSETS = ROOT / "docs" / "assets"
DEFAULT_OUTPUT = ROOT / "reports" / "presentation" / "Customer_Intelligence_ML_Pipeline.pptx"

SLIDE_W = 13.333
SLIDE_H = 7.5

NAVY = "071628"
NAVY_2 = "0B1D33"
CARD = "102842"
CARD_2 = "132F4D"
CYAN = "55D7E3"
BLUE = "4E8DFF"
VIOLET = "947CFF"
ORANGE = "FF9C5B"
WHITE = "F7FAFF"
MUTED = "AFC0D1"
GRID = "22435F"
GREEN = "4ED6A6"
RED = "FF6E78"


def rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


def add_full_bleed(slide, image_path: Path) -> None:
    slide.shapes.add_picture(str(image_path), 0, 0, width=Inches(SLIDE_W), height=Inches(SLIDE_H))


def add_rect(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    fill: str = CARD,
    line: str | None = GRID,
    radius: bool = True,
    line_width: float = 1.0,
):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if line:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(line_width)
    else:
        shape.line.fill.background()
    return shape


def add_text(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    size: float = 16,
    color: str = WHITE,
    bold: bool = False,
    font: str = "Aptos",
    align: PP_ALIGN = PP_ALIGN.LEFT,
    valign: MSO_ANCHOR = MSO_ANCHOR.TOP,
    margin: float = 0.03,
    line_spacing: float = 1.0,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = align
    paragraph.line_spacing = line_spacing
    for run in paragraph.runs:
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = rgb(color)
    return box


def add_bullets(
    slide,
    bullets: list[str],
    x: float,
    y: float,
    w: float,
    h: float,
    size: float = 15,
    color: str = WHITE,
    bullet_color: str = CYAN,
    spacing: float = 6,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.04)
    frame.margin_right = Inches(0.04)
    frame.margin_top = Inches(0.02)
    frame.margin_bottom = Inches(0.02)
    for index, item in enumerate(bullets):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = f"•  {item}"
        paragraph.space_after = Pt(spacing)
        paragraph.line_spacing = 1.05
        for run in paragraph.runs:
            run.font.name = "Aptos"
            run.font.size = Pt(size)
            run.font.color.rgb = rgb(color)
        if paragraph.runs:
            paragraph.runs[0].font.color.rgb = rgb(bullet_color)
    return box


def add_background(slide) -> None:
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = rgb(NAVY)
    add_rect(slide, 0.0, 0.0, SLIDE_W, 0.10, fill=CYAN, line=None, radius=False)
    add_rect(slide, 0.0, 7.42, SLIDE_W, 0.08, fill=VIOLET, line=None, radius=False)


def add_header(slide, title: str, section: str, slide_number: int) -> None:
    add_text(slide, section.upper(), 0.58, 0.22, 6.0, 0.22, size=9.5, color=CYAN, bold=True)
    add_text(slide, title, 0.58, 0.48, 11.9, 0.52, size=28, color=WHITE, bold=True)
    add_text(
        slide,
        f"{slide_number:02d}",
        12.22,
        0.25,
        0.52,
        0.27,
        size=11,
        color=MUTED,
        bold=True,
        align=PP_ALIGN.RIGHT,
    )


def add_footer(slide) -> None:
    add_text(
        slide,
        "github.com/Awais-Asghar/customer-intelligence-ml",
        0.58,
        7.14,
        6.2,
        0.18,
        size=8.5,
        color=MUTED,
    )


def add_notes(slide, text: str) -> None:
    slide.notes_slide.notes_text_frame.text = text


def add_image_contain(
    slide,
    image_path: Path,
    x: float,
    y: float,
    w: float,
    h: float,
    border: str = GRID,
    background: str = WHITE,
):
    add_rect(slide, x, y, w, h, fill=background, line=border, radius=True, line_width=1.1)
    with Image.open(image_path) as image:
        image_w, image_h = image.size
    scale = min((w - 0.12) / image_w, (h - 0.12) / image_h)
    draw_w = image_w * scale
    draw_h = image_h * scale
    draw_x = x + (w - draw_w) / 2
    draw_y = y + (h - draw_h) / 2
    picture = slide.shapes.add_picture(
        str(image_path),
        Inches(draw_x),
        Inches(draw_y),
        width=Inches(draw_w),
        height=Inches(draw_h),
    )
    return picture


def add_stat_card(
    slide,
    value: str,
    label: str,
    x: float,
    y: float,
    w: float,
    h: float,
    accent: str = CYAN,
    value_size: float = 23,
):
    add_rect(slide, x, y, w, h, fill=CARD, line=GRID)
    add_rect(slide, x, y, 0.06, h, fill=accent, line=None, radius=False)
    add_text(slide, value, x + 0.18, y + 0.12, w - 0.30, h * 0.46, size=value_size, color=accent, bold=True)
    add_text(slide, label, x + 0.18, y + h * 0.57, w - 0.30, h * 0.28, size=10.5, color=MUTED, bold=True)


def add_method_card(slide, x: float, y: float, w: float, accent: str, number: str, title: str, question: str, role: str):
    add_rect(slide, x, y, w, 4.58, fill=CARD, line=accent, line_width=1.4)
    add_text(slide, number, x + 0.24, y + 0.18, 0.55, 0.42, size=20, color=accent, bold=True)
    add_text(slide, title, x + 0.24, y + 0.78, w - 0.48, 0.40, size=19, color=WHITE, bold=True)
    add_text(slide, question, x + 0.24, y + 1.36, w - 0.48, 1.05, size=16, color=WHITE, bold=False)
    add_text(slide, "ROLE", x + 0.24, y + 2.75, w - 0.48, 0.22, size=9.5, color=accent, bold=True)
    add_text(slide, role, x + 0.24, y + 3.08, w - 0.48, 1.02, size=13.5, color=MUTED)


def add_metric_row(slide, items: list[tuple[str, str, str]], y: float) -> None:
    count = len(items)
    gap = 0.12
    total_w = 12.17
    card_w = (total_w - (count - 1) * gap) / count
    x = 0.58
    for value, label, accent in items:
        add_stat_card(slide, value, label, x, y, card_w, 0.90, accent=accent, value_size=18)
        x += card_w + gap


def build_deck(output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    presentation = Presentation()
    presentation.slide_width = Inches(SLIDE_W)
    presentation.slide_height = Inches(SLIDE_H)
    blank = presentation.slide_layouts[6]

    presentation.core_properties.title = "Customer Intelligence ML Pipeline"
    presentation.core_properties.subject = "Machine learning portfolio presentation"
    presentation.core_properties.author = "Awais Asghar"
    presentation.core_properties.keywords = "PCA, K-Means, DBSCAN, gradient boosting, machine learning"

    # Slide 1
    slide = presentation.slides.add_slide(blank)
    add_full_bleed(slide, ASSETS / "customer-intelligence-banner.png")
    add_rect(slide, 0.0, 5.05, SLIDE_W, 2.45, fill=NAVY, line=None, radius=False)
    add_rect(slide, 0.0, 5.05, 0.10, 2.45, fill=ORANGE, line=None, radius=False)
    add_text(slide, "CUSTOMER INTELLIGENCE", 0.66, 5.34, 8.8, 0.64, size=34, color=WHITE, bold=True)
    add_text(
        slide,
        "Segmentation, anomaly exploration, dimensionality reduction, and calibrated default-risk modeling",
        0.68,
        6.04,
        10.9,
        0.56,
        size=17,
        color=MUTED,
    )
    add_text(slide, "ML PORTFOLIO CASE STUDY", 0.68, 6.73, 3.2, 0.24, size=10, color=CYAN, bold=True)
    add_text(
        slide,
        "github.com/Awais-Asghar/customer-intelligence-ml",
        8.22,
        6.72,
        4.48,
        0.22,
        size=9.5,
        color=MUTED,
        align=PP_ALIGN.RIGHT,
    )
    add_notes(slide, "Open by explaining that this project demonstrates complete machine-learning reasoning, not only model training. The four methods answer separate questions and their outputs must not be mixed.")

    # Slide 2
    slide = presentation.slides.add_slide(blank)
    add_background(slide)
    add_header(slide, "Four questions, four methods", "Business framing", 2)
    cards = [
        (0.58, CYAN, "01", "PCA", "What structure explains most variation?", "Compression and visualization"),
        (3.73, BLUE, "02", "K-Means", "Which broad behaviors repeat?", "Descriptive segmentation"),
        (6.88, VIOLET, "03", "DBSCAN", "Which records are locally sparse?", "Anomaly exploration"),
        (10.03, ORANGE, "04", "Boosting", "How uncertain is historical default?", "Calibrated risk estimation"),
    ]
    for x, accent, number, title, question, role in cards:
        add_method_card(slide, x, 1.38, 2.72, accent, number, title, question, role)
    add_rect(slide, 1.42, 6.23, 10.50, 0.58, fill=CARD_2, line=GRID)
    add_text(
        slide,
        "Clusters are not risk classes. Density noise is not fraud. Predictions are probabilities, not decisions.",
        1.65,
        6.38,
        10.05,
        0.24,
        size=14,
        color=WHITE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_footer(slide)
    add_notes(slide, "Explain the boundary between the methods. PCA compresses data, K-Means describes broad groups, DBSCAN finds sparse behavior, and gradient boosting estimates a historical outcome probability.")

    # Slide 3
    slide = presentation.slides.add_slide(blank)
    add_background(slide)
    add_header(slide, "A realistic, imbalanced credit dataset", "Data", 3)
    add_image_contain(slide, FIGURES / "class-distribution.png", 0.58, 1.33, 7.10, 4.40)
    add_stat_card(slide, "30,000", "CUSTOMER RECORDS", 8.05, 1.34, 2.15, 1.18, CYAN)
    add_stat_card(slide, "23", "RAW EXPLANATORY FEATURES", 10.47, 1.34, 2.28, 1.18, BLUE)
    add_stat_card(slide, "22.12%", "HISTORICAL DEFAULTS", 8.05, 2.74, 2.15, 1.18, ORANGE)
    add_stat_card(slide, "0", "MISSING VALUES", 10.47, 2.74, 2.28, 1.18, GREEN)
    add_rect(slide, 8.05, 4.19, 4.70, 1.54, fill=CARD, line=GRID)
    add_text(slide, "WHY ACCURACY IS NOT ENOUGH", 8.28, 4.42, 4.20, 0.24, size=10, color=ORANGE, bold=True)
    add_text(
        slide,
        "A model that predicts only the majority class would appear accurate while missing every historical default.",
        8.28,
        4.78,
        4.12,
        0.70,
        size=14.5,
        color=WHITE,
    )
    add_rect(slide, 0.58, 6.03, 12.17, 0.72, fill=CARD_2, line=GRID)
    add_text(slide, "Source", 0.83, 6.22, 0.68, 0.22, size=10, color=CYAN, bold=True)
    add_text(slide, "UCI Default of Credit Card Clients, dataset 350, CC BY 4.0", 1.52, 6.17, 6.4, 0.28, size=13, color=WHITE)
    add_text(slide, "Target prevalence", 8.75, 6.22, 1.42, 0.22, size=10, color=ORANGE, bold=True)
    add_text(slide, "6,636 positive records", 10.17, 6.17, 2.28, 0.28, size=13, color=WHITE)
    add_footer(slide)
    add_notes(slide, "The target is imbalanced but not extremely rare. Use balanced accuracy, precision, recall, F1, PR-AUC, and ROC-AUC alongside overall accuracy.")

    # Slide 4
    slide = presentation.slides.add_slide(blank)
    add_background(slide)
    add_header(slide, "The test set stays closed until the end", "Leakage control", 4)
    flow = [
        (0.67, 3.83, CYAN, "TRAIN", "24,000 records | 80%", "Candidate selection\nThree-fold stratified CV"),
        (4.75, 3.83, VIOLET, "VALIDATION", "3,000 records | 10%", "Threshold selection\nNo candidate tuning"),
        (8.83, 3.83, ORANGE, "TEST", "3,000 records | 10%", "Opened once\nFinal evaluation only"),
    ]
    for x, w, accent, name, count, purpose in flow:
        add_rect(slide, x, 1.48, w, 2.10, fill=CARD, line=accent, line_width=1.5)
        add_text(slide, name, x + 0.24, 1.72, w - 0.48, 0.32, size=17, color=accent, bold=True)
        add_text(slide, count, x + 0.24, 2.18, w - 0.48, 0.30, size=14, color=WHITE, bold=True)
        add_text(slide, purpose, x + 0.24, 2.70, w - 0.48, 0.60, size=13, color=MUTED)
    add_rect(slide, 0.67, 4.02, 12.08, 2.33, fill=CARD_2, line=GRID)
    add_text(slide, "AUDIT RESULTS", 0.95, 4.29, 2.0, 0.24, size=11, color=CYAN, bold=True)
    audit_items = [
        ("0", "identifier overlap", GREEN),
        ("0.0742", "maximum absolute SMD", CYAN),
        ("10", "test rows matching a dev profile", VIOLET),
        ("21", "profiles with conflicting labels", ORANGE),
    ]
    x = 0.95
    for value, label, accent in audit_items:
        add_stat_card(slide, value, label.upper(), x, 4.77, 2.65, 1.18, accent=accent, value_size=22)
        x += 2.87
    add_footer(slide)
    add_notes(slide, "Candidate selection uses only training folds. Validation is reserved for threshold selection. The test set is opened once after the model and threshold are frozen. The conflicting labels explain why perfect classification is not credible.")

    # Slide 5
    slide = presentation.slides.add_slide(blank)
    add_background(slide)
    add_header(slide, "PCA retains signal without inventing certainty", "Dimensionality reduction", 5)
    add_image_contain(slide, FIGURES / "pca-preprocessing-comparison.png", 0.58, 1.29, 5.34, 3.38)
    add_image_contain(slide, FIGURES / "pca-diagnostic-projection.png", 6.10, 1.29, 6.65, 3.38)
    add_metric_row(
        slide,
        [
            ("19", "BEHAVIORAL INPUTS", CYAN),
            ("13", "COMPONENTS RETAINED", BLUE),
            ("95.63%", "CUMULATIVE VARIANCE", VIOLET),
            ("59.29%", "FIRST TWO COMPONENTS", ORANGE),
        ],
        4.94,
    )
    add_rect(slide, 0.58, 6.03, 12.17, 0.72, fill=CARD_2, line=GRID)
    add_text(
        slide,
        "The 13-component representation supports PCA and clustering analysis. The supervised model uses 54 engineered behavioral features.",
        0.84,
        6.20,
        11.64,
        0.30,
        size=14,
        color=WHITE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_footer(slide)
    add_notes(slide, "Preprocessing and PCA were fitted on training data only. Two components are useful for visualization, while the retained representation supports unsupervised analysis. The supervised classifier does not use the PCA representation.")

    # Slide 6
    slide = presentation.slides.add_slide(blank)
    add_background(slide)
    add_header(slide, "Two stable groups, with meaningful overlap", "K-Means segmentation", 6)
    add_image_contain(slide, FIGURES / "kmeans-model-selection.png", 0.58, 1.28, 6.08, 4.64)
    add_image_contain(slide, FIGURES / "kmeans-cluster-projection.png", 6.84, 1.28, 5.91, 4.64)
    add_metric_row(
        slide,
        [
            ("2", "SELECTED CLUSTERS", CYAN),
            ("0.2638", "SILHOUETTE SCORE", BLUE),
            ("0.9960", "STABILITY ARI", VIOLET),
            ("24.66%", "SMALLEST CLUSTER", ORANGE),
        ],
        6.07,
    )
    add_footer(slide)
    add_notes(slide, "The near-perfect stability ARI means repeated initializations find almost the same assignments. The moderate silhouette score shows overlap, so the groups should be presented as broad behavioral segments, not natural risk classes.")

    # Slide 7
    slide = presentation.slides.add_slide(blank)
    add_background(slide)
    add_header(slide, "DBSCAN reveals sparse records, not new segments", "Anomaly exploration", 7)
    add_image_contain(slide, FIGURES / "dbscan-distance-curve.png", 0.58, 1.30, 5.18, 2.65)
    add_rect(slide, 0.58, 4.17, 5.18, 2.23, fill=CARD, line=GRID)
    add_text(slide, "SELECTED CONFIGURATION", 0.84, 4.43, 3.0, 0.24, size=10, color=VIOLET, bold=True)
    add_bullets(
        slide,
        [
            "5 whitened PCA components",
            "eps 1.1526 and min_samples 10",
            "One substantial core group",
            "2.08% density-noise records",
        ],
        0.84,
        4.82,
        4.55,
        1.34,
        size=14,
        bullet_color=VIOLET,
        spacing=3,
    )
    add_image_contain(slide, FIGURES / "dbscan-anomaly-projection.png", 5.96, 1.30, 6.79, 5.10)
    add_rect(slide, 0.58, 6.56, 12.17, 0.35, fill=ORANGE, line=None, radius=True)
    add_text(
        slide,
        "Interpretation: unusual multivariate behavior, not fraud, wrongdoing, or a guaranteed future outcome.",
        0.78,
        6.63,
        11.77,
        0.18,
        size=11.5,
        color=NAVY,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_footer(slide)
    add_notes(slide, "No tested configuration produced multiple credible density clusters. Keeping the method as anomaly exploration is a scientifically useful negative result. Never call the noise points fraud.")

    # Slide 8
    slide = presentation.slides.add_slide(blank)
    add_background(slide)
    add_header(slide, "A supervised pipeline designed for honest comparison", "Gradient boosting", 8)
    stages = [
        ("01", "Behavioral inputs", "19 non-demographic variables", CYAN),
        ("02", "Feature engineering", "54 deterministic features", BLUE),
        ("03", "Candidate search", "10 HGB configurations", VIOLET),
        ("04", "Selection", "Three-fold stratified CV", ORANGE),
    ]
    x = 0.58
    for index, (number, title, detail, accent) in enumerate(stages):
        add_rect(slide, x, 1.50, 2.72, 2.22, fill=CARD, line=accent, line_width=1.4)
        add_text(slide, number, x + 0.22, 1.72, 0.52, 0.30, size=15, color=accent, bold=True)
        add_text(slide, title, x + 0.22, 2.20, 2.24, 0.34, size=17, color=WHITE, bold=True)
        add_text(slide, detail, x + 0.22, 2.78, 2.24, 0.54, size=13.5, color=MUTED)
        if index < len(stages) - 1:
            chevron = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(x + 2.81), Inches(2.29), Inches(0.33), Inches(0.52))
            chevron.fill.solid()
            chevron.fill.fore_color.rgb = rgb(GRID)
            chevron.line.fill.background()
        x += 3.08
    add_rect(slide, 0.58, 4.11, 12.17, 1.10, fill=CARD_2, line=GRID)
    add_text(slide, "MODEL SELECTION OBJECTIVE", 0.85, 4.33, 2.35, 0.20, size=9.5, color=CYAN, bold=True)
    add_text(slide, "PR-AUC", 0.85, 4.62, 1.35, 0.34, size=22, color=WHITE, bold=True)
    add_text(slide, "Selected because the positive class is 22.12% of the data", 2.34, 4.62, 5.45, 0.35, size=14, color=MUTED)
    add_text(slide, "Demographics excluded from prediction", 8.25, 4.58, 4.08, 0.34, size=14, color=ORANGE, bold=True, align=PP_ALIGN.RIGHT)
    add_metric_row(
        slide,
        [
            ("0.05", "LEARNING RATE", CYAN),
            ("15", "MAX LEAF NODES", BLUE),
            ("600", "MAX ITERATIONS", VIOLET),
            ("30", "EARLY-STOP PATIENCE", ORANGE),
        ],
        5.56,
    )
    add_footer(slide)
    add_notes(slide, "The feature contract is deterministic. Demographic fields are excluded from prediction and reserved for aggregate auditing. Histogram gradient boosting uses gradient-based tree updates, not Adam or SGD.")

    # Slide 9
    slide = presentation.slides.add_slide(blank)
    add_background(slide)
    add_header(slide, "Early stopping controls unnecessary boosting", "Training diagnostics", 9)
    add_image_contain(slide, FIGURES / "training-validation-loss.png", 0.58, 1.25, 12.17, 3.26)
    add_metric_row(
        slide,
        [
            ("81", "FOLD 1 BEST ITERATION", CYAN),
            ("72", "FOLD 2 BEST ITERATION", BLUE),
            ("138", "FOLD 3 BEST ITERATION", VIOLET),
            ("600", "MAXIMUM CAP", ORANGE),
        ],
        4.73,
    )
    add_rect(slide, 0.58, 5.89, 12.17, 0.86, fill=CARD_2, line=GRID)
    add_text(slide, "Final calibrated estimators", 0.83, 6.11, 2.12, 0.23, size=10, color=CYAN, bold=True)
    add_text(slide, "116, 131, and 106 iterations", 2.88, 6.06, 3.44, 0.29, size=16, color=WHITE, bold=True)
    add_text(slide, "No fold reached the cap", 9.42, 6.10, 3.00, 0.25, size=13, color=GREEN, bold=True, align=PP_ALIGN.RIGHT)
    add_footer(slide)
    add_notes(slide, "The real notebook curves are shown here. Validation loss plateaus while training loss keeps falling, especially in fold three. Raising the 600-iteration cap would not improve this run.")

    # Slide 10
    slide = presentation.slides.add_slide(blank)
    add_background(slide)
    add_header(slide, "Threshold selection is an operating decision", "Calibration", 10)
    add_image_contain(slide, FIGURES / "threshold-and-calibration.png", 0.58, 1.26, 12.17, 3.71)
    add_stat_card(slide, "0.2916", "VALIDATION-SELECTED THRESHOLD", 0.58, 5.22, 3.12, 1.12, ORANGE, value_size=24)
    add_rect(slide, 3.93, 5.22, 4.20, 1.12, fill=CARD, line=GRID)
    add_text(slide, "GAINS", 4.17, 5.42, 0.65, 0.22, size=10, color=GREEN, bold=True)
    add_text(slide, "Recall, balanced accuracy, and F1", 4.17, 5.75, 3.70, 0.28, size=14.5, color=WHITE, bold=True)
    add_rect(slide, 8.36, 5.22, 4.39, 1.12, fill=CARD, line=GRID)
    add_text(slide, "TRADEOFF", 8.60, 5.42, 0.90, 0.22, size=10, color=ORANGE, bold=True)
    add_text(slide, "Lower precision and overall accuracy", 8.60, 5.75, 3.88, 0.28, size=14.5, color=WHITE, bold=True)
    add_rect(slide, 0.58, 6.55, 12.17, 0.36, fill=CARD_2, line=GRID)
    add_text(
        slide,
        "ROC-AUC and PR-AUC do not change when only the decision threshold changes.",
        0.84,
        6.62,
        11.64,
        0.18,
        size=11.5,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )
    add_footer(slide)
    add_notes(slide, "The chart is the actual validation result. The selected threshold maximizes validation F1. Lowering the threshold flags more positives, improving recall while increasing false positives.")

    # Slide 11
    slide = presentation.slides.add_slide(blank)
    add_background(slide)
    add_header(slide, "The final test result is useful and honest", "Final evaluation", 11)
    add_image_contain(slide, FIGURES / "test-confusion-matrix-and-roc.png", 0.58, 1.24, 12.17, 4.30)
    add_metric_row(
        slide,
        [
            ("80.47%", "ACCURACY", CYAN),
            ("71.02%", "BALANCED ACCURACY", BLUE),
            ("56.09%", "PRECISION", VIOLET),
            ("54.07%", "RECALL", ORANGE),
            ("55.06%", "F1 SCORE", GREEN),
        ],
        5.74,
    )
    add_footer(slide)
    add_notes(slide, "At the selected threshold, the confusion matrix contains 2,055 true negatives, 281 false positives, 305 false negatives, and 359 true positives. Compare this operating point with threshold 0.50 and explain the tradeoff.")

    # Slide 12
    slide = presentation.slides.add_slide(blank)
    add_background(slide)
    add_header(slide, "What the project proves, and what it does not", "Responsible ML", 12)
    add_rect(slide, 0.58, 1.28, 5.93, 4.70, fill=CARD, line=CYAN, line_width=1.3)
    add_text(slide, "LESSONS", 0.87, 1.57, 1.20, 0.24, size=11, color=CYAN, bold=True)
    add_bullets(
        slide,
        [
            "Prevent leakage before optimizing metrics",
            "Use separate methods for separate questions",
            "Tune the threshold on validation data only",
            "Treat negative results as useful evidence",
            "Communicate uncertainty and operating tradeoffs",
        ],
        0.87,
        2.06,
        5.25,
        2.70,
        size=16,
        bullet_color=CYAN,
        spacing=9,
    )
    add_rect(slide, 6.82, 1.28, 5.93, 4.70, fill=CARD, line=ORANGE, line_width=1.3)
    add_text(slide, "LIMITS", 7.11, 1.57, 1.20, 0.24, size=11, color=ORANGE, bold=True)
    add_bullets(
        slide,
        [
            "Conflicting labels prevent deterministic prediction",
            "Rare categories increase uncertainty",
            "Clusters are descriptive, not causal",
            "DBSCAN noise is not evidence of fraud",
            "Production use requires governance and monitoring",
        ],
        7.11,
        2.06,
        5.25,
        2.70,
        size=16,
        bullet_color=ORANGE,
        spacing=9,
    )
    add_rect(slide, 0.58, 6.22, 12.17, 0.68, fill=CARD_2, line=GRID)
    add_text(slide, "REPOSITORY", 0.84, 6.43, 1.10, 0.20, size=9.5, color=CYAN, bold=True)
    add_text(slide, "github.com/Awais-Asghar/customer-intelligence-ml", 1.91, 6.36, 5.64, 0.29, size=15, color=WHITE, bold=True)
    add_text(slide, "Notebook  |  Tests  |  Figures  |  Reproducibility", 7.66, 6.39, 4.77, 0.24, size=12.5, color=MUTED, align=PP_ALIGN.RIGHT)
    add_footer(slide)
    add_notes(slide, "Close with the main engineering lesson: trustworthy ML is not a contest for the largest metric. The project demonstrates leakage control, reproducibility, method selection, calibration, and honest communication of limitations.")

    presentation.save(output_path)
    print(f"Saved {len(presentation.slides)} slides to {output_path}")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    return parser.parse_args()


if __name__ == "__main__":
    arguments = parse_args()
    build_deck(arguments.output.resolve())
